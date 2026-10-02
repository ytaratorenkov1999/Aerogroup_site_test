import json
from unittest.mock import patch, MagicMock
from datetime import datetime
from django.test import TestCase
from django.urls import reverse
from statistic.models import HelpdeskStatistic, HelpdeskAdminStat, HelpdeskCategoryAfl, HelpdeskCategoryAkr, HelpdeskTicket
from statistic.helpdesk_api import Helpdesk
from support.tests_helpers import make_user, LoginMixin
from support.roles import ROLE_READER, ROLE_ADMIN
import io
import openpyxl
from statistic.helpdesk_api import HelpdeskSaver, first_answer_minutes
from statistic.export_to_excel import REPORTS

class DashboardAccessTests(LoginMixin, TestCase):

    def test_anonymous_redirected(self):
        response = self.client.get(reverse('stats_hd'))
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response['Location'])

    def test_logged_in_user_sees_dashboard(self):
        user = make_user('stat1', ROLE_READER)
        self.login(user)
        response = self.client.get(reverse('stats_hd'))
        self.assertEqual(response.status_code, 200)

    def test_dashboard_with_no_data_renders(self):
        """Страница должна открываться даже без записей в БД."""
        user = make_user('stat2', ROLE_READER)
        self.login(user)
        response = self.client.get(reverse('stats_hd') + '?period=2025-01')
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.context.get('stat'))

    def test_dashboard_with_existing_period(self):
        stat = HelpdeskStatistic.objects.create(
            period='2025-05',
            total_ticket=10,
            request_count_afl=6,
            request_count_akr=4,
            count_service_ife=2,
            count_service_rportal=1,
            redirect_second_line=15.0,
            avg_time_first_answer=3.5,
        )
        user = make_user('stat3', ROLE_READER)
        self.login(user)
        response = self.client.get(reverse('stats_hd') + '?period=2025-05')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['stat'].id, stat.id)



class UpdateStatisticTests(LoginMixin, TestCase):

    def setUp(self):
        self.user = make_user('stat_upd', ROLE_ADMIN)
        self.login(self.user)
        self.url = reverse('stats_hd_update')

    def test_get_not_allowed(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 405)

    def test_invalid_period_format_returns_400(self):
        response = self.client.post(self.url, {'period': 'not-a-date'})
        self.assertEqual(response.status_code, 400)
        data = json.loads(response.content)
        self.assertFalse(data['ok'])

    def test_future_period_returns_400(self):
        response = self.client.post(self.url, {'period': '2099-01'})
        self.assertEqual(response.status_code, 400)
        data = json.loads(response.content)
        self.assertFalse(data['ok'])

    @patch('statistic.views.Helpdesk')
    @patch('statistic.views.HelpdeskSaver')
    def test_successful_update_returns_ok(self, mock_saver, mock_helpdesk_cls):
        mock_instance = MagicMock()
        mock_instance.get_all_statistic.return_value = {'total': 5}
        mock_helpdesk_cls.return_value = mock_instance

        now = datetime.now()
        period = now.strftime('%Y-%m')
        response = self.client.post(self.url, {'period': period})
        self.assertEqual(response.status_code, 200)
        data = json.loads(response.content)
        self.assertTrue(data['ok'])

    def test_anonymous_cannot_update(self):
        self.client.logout()
        response = self.client.post(self.url, {'period': '2025-01'})
        self.assertEqual(response.status_code, 302)




class ExportExcelTests(LoginMixin, TestCase):

    def setUp(self):
        self.user = make_user('stat_exp', ROLE_READER)
        self.login(self.user)

    def test_export_without_data_returns_404_or_redirect(self):
        response = self.client.get(reverse('stats_hd_export') + '?period=2099-12')
        # Нет данных — ожидаем либо 404, либо редирект обратно
        self.assertIn(response.status_code, [302, 404])

    def test_export_with_data_returns_xlsx(self):
        stat = HelpdeskStatistic.objects.create(
            period='2025-06',
            total_ticket=20,
            request_count_afl=12,
            request_count_akr=8,
            count_service_ife=5,
            count_service_rportal=3,
            redirect_second_line=10.0,
            avg_time_first_answer=None,
        )
        HelpdeskAdminStat.objects.create(statistic=stat, name='Иванов', count=10)
        HelpdeskCategoryAfl.objects.create(statistic=stat, category='IFE', count=5, ticket_ids=[1, 2])
        HelpdeskCategoryAkr.objects.create(statistic=stat, category='Portal', count=3, ticket_ids=[3])

        response = self.client.get(reverse('stats_hd_export') + '?period=2025-06')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_anonymous_redirected_from_export(self):
        self.client.logout()
        response = self.client.get(reverse('stats_hd_export'))
        self.assertEqual(response.status_code, 302)



class HelpdeskValidatePeriodTests(TestCase):

    @patch('statistic.helpdesk_api.config')
    def _make_helpdesk(self, mock_config, period):
        mock_config.return_value = 'dummy'
        return Helpdesk(period=period)

    def test_valid_past_period_does_not_raise(self):
        with patch('statistic.helpdesk_api.config', return_value='dummy'):
            # Не должно бросать исключение
            hd = Helpdesk.__new__(Helpdesk)
            hd.period = '2024-01'
            hd.validate_period()

    def test_future_period_raises_value_error(self):
        with patch('statistic.helpdesk_api.config', return_value='dummy'):
            hd = Helpdesk.__new__(Helpdesk)
            hd.period = '2099-12'
            with self.assertRaises(ValueError):
                hd.validate_period()

    def test_invalid_format_raises_value_error(self):
        with patch('statistic.helpdesk_api.config', return_value='dummy'):
            hd = Helpdesk.__new__(Helpdesk)
            hd.period = 'invalid'
            with self.assertRaises(ValueError):
                hd.validate_period()

    def test_wrong_separator_raises_value_error(self):
        with patch('statistic.helpdesk_api.config', return_value='dummy'):
            hd = Helpdesk.__new__(Helpdesk)
            hd.period = '2024/01'
            with self.assertRaises(ValueError):
                hd.validate_period()


class HelpdeskStatisticModelTests(TestCase):

    def _make_stat(self, period='2025-03'):
        return HelpdeskStatistic.objects.create(
            period=period,
            total_ticket=30,
            request_count_afl=18,
            request_count_akr=12,
            count_service_ife=7,
            count_service_rportal=4,
            redirect_second_line=20.0,
            avg_time_first_answer=5.0,
        )

    def test_str_contains_period(self):
        stat = self._make_stat('2025-03')
        self.assertIn('2025-03', str(stat))

    def test_ordering_by_created_at_desc(self):
        s1 = self._make_stat('2025-01')
        s2 = self._make_stat('2025-02')
        qs = list(HelpdeskStatistic.objects.all())
        # Последний созданный идёт первым
        self.assertEqual(qs[0].id, s2.id)

    def test_admin_stat_str(self):
        stat = self._make_stat('2025-04')
        admin = HelpdeskAdminStat.objects.create(statistic=stat, name='Петров', count=7)
        self.assertIn('Петров', str(admin))
        self.assertIn('7', str(admin))

    def test_category_afl_str(self):
        stat = self._make_stat('2025-05')
        cat = HelpdeskCategoryAfl.objects.create(statistic=stat, category='IFE issues', count=3)
        self.assertIn('IFE issues', str(cat))

    def test_ticket_ids_default_empty_list(self):
        stat = self._make_stat('2025-06')
        cat = HelpdeskCategoryAkr.objects.create(statistic=stat, category='Portal', count=2)
        self.assertEqual(cat.ticket_ids, [])

    def test_avg_time_can_be_null(self):
        stat = HelpdeskStatistic.objects.create(
            period='2025-07',
            total_ticket=0, request_count_afl=0, request_count_akr=0,
            count_service_ife=0, count_service_rportal=0,
            redirect_second_line=0.0, avg_time_first_answer=None,
        )
        self.assertIsNone(stat.avg_time_first_answer)


ENV = {
    'SERVICE_NAME_AFL': 'AFL CrewTab',
    'SERVICE_NAME_AKR': 'SDM Crew Tablet',
    'SECOND_LINE': 'Вторая линия поддержки',
}


def _resp(rows):
    r = MagicMock()
    r.json.return_value = rows
    return r


class SecondLineTests(TestCase):

    def setUp(self):
        patcher = patch('statistic.helpdesk_api.config', side_effect=lambda k, *a, **kw: ENV.get(k, kw.get('default')))
        patcher.start()
        self.addCleanup(patcher.stop)
        self.hd = Helpdesk.__new__(Helpdesk)

    def test_counts_only_crewtab_services(self):
        afl = _resp([
            {'service_name': 'AFL CrewTab', 'gfullname': 'Вторая линия поддержки'},
            {'service_name': 'AFL CrewTab', 'gfullname': 'Первая линия'},
            {'service_name': 'IFE.AFL',     'gfullname': 'Вторая линия поддержки'},  # не считается
        ])
        akr = _resp([
            {'service_name': 'SDM Crew Tablet', 'gfullname': 'Вторая линия поддержки'},
            {'service_name': 'rportal.aero',    'gfullname': 'Вторая линия поддержки'},  # не считается
        ])
        count, percent = self.hd.second_line_stats(afl, akr)
        self.assertEqual(count, 2)
        self.assertEqual(percent, 66.7)   # 2 из 3 CrewTab-заявок

    def test_zero_tickets_no_division_error(self):
        self.assertEqual(self.hd.second_line_stats(_resp([]), _resp([])), (0, 0.0))

    def test_percent_never_exceeds_100(self):
        rows = [{'service_name': 'IFE.AFL', 'gfullname': 'Вторая линия поддержки'}] * 10
        rows.append({'service_name': 'AFL CrewTab', 'gfullname': 'Вторая линия поддержки'})
        count, percent = self.hd.second_line_stats(_resp(rows), _resp([]))
        self.assertEqual((count, percent), (1, 100.0))




REPORT_ENV = {
    'SERVICE_NAME_AFL': 'AFL CrewTab',
    'SERVICE_NAME_AKR': 'SDM Crew Tablet',
    'SERVICE_NAME_IFE_AFL': 'IFE.AFL',
    'SERVICE_NAME_MEDIA_AKR': 'rportal.aero',
    'SECOND_LINE': 'Вторая линия поддержки',
}


def _fake_config(key, *args, **kwargs):
    if key in REPORT_ENV:
        return REPORT_ENV[key]
    if 'default' in kwargs:
        return kwargs['default']
    raise KeyError(key)


def _row(tid, service, group='Первая линия', executor='', category='Отчёты',
         date='01.09.2026 10:00', start='01.09.2026 10:30'):
    return {'id': tid, 'service_name': service, 'gfullname': group, 'mfullname': executor,
            'ZayavCategory_id': category, 'Date': date, 'fStartTime': start}


AFL_ROWS = [
    _row(1, 'AFL CrewTab', executor='Павлов Андрей'),
    _row(2, 'AFL CrewTab', group='Вторая линия поддержки', executor='Тараторенков Юрий',
         category='E-mail ticket', start=''),
    _row(3, 'IFE.AFL', executor='Павлов Андрей'),                 # сервис, не Crew Tablet
]
AKR_ROWS = [
    _row(10, 'SDM Crew Tablet', executor='Романов Владислав', category='Экипаж',
         date='02.09.2026 09:00', start='02.09.2026 09:10'),
    _row(11, 'rportal.aero', group='Вторая линия поддержки'),     # не входит во 2 линию
]


class TicketDetailsTests(LoginMixin, TestCase):

    def setUp(self):
        patcher = patch('statistic.models.config', side_effect=_fake_config)
        patcher.start()
        self.addCleanup(patcher.stop)
        patcher2 = patch('statistic.helpdesk_api.config', side_effect=_fake_config)
        patcher2.start()
        self.addCleanup(patcher2.stop)

        afl, akr = _resp(AFL_ROWS), _resp(AKR_ROWS)
        hd = Helpdesk.__new__(Helpdesk)
        count, percent = hd.second_line_stats(afl, akr)
        cats_afl, ids_afl = Helpdesk.count_categories_afl(afl)
        cats_akr, ids_akr = Helpdesk.count_categories_akr(akr)
        data = {
            'total_ticket': Helpdesk.total_tickets(afl, akr),
            'request_count_afl': 2, 'request_count_akr': 1,
            'count_service_ife': 1, 'count_service_rportal': 1,
            'redirect_second_line': percent, 'redirect_second_line_count': count,
            'avg_time_first_answer': hd.avg_first_answers_user(afl, akr),
            'count_user_request_total': Helpdesk.count_user_request_total(
                afl, akr, ['Павлов Андрей', 'Тараторенков Юрий', 'Романов Владислав']),
            'count_categories_afl': cats_afl, 'ids_categories_afl': ids_afl,
            'count_categories_akr': cats_akr, 'ids_categories_akr': ids_akr,
            'tickets': Helpdesk.collect_tickets(afl, akr),
        }
        self.data = data
        HelpdeskSaver.save(data, '2026-09')
        self.stat = HelpdeskStatistic.objects.get(period='2026-09')
        self.login(make_user('stat_rep', ROLE_READER))

    def _download(self, report):
        response = self.client.get(reverse('stats_hd_export') + f'?period=2026-09&report={report}')
        self.assertEqual(response.status_code, 200, report)
        return openpyxl.load_workbook(io.BytesIO(response.content))

    @staticmethod
    def _ids(ws):
        """ID заявок из колонки «Ссылка на заявку» (D — таблица начинается с C3)"""
        return [int(c.value.rsplit('/', 1)[1]) for c in ws['D'] if isinstance(c.value, str) and '/request/' in c.value]


    def test_admin_stat_counts_only_crew_tablet(self):
        self.assertEqual(self.data['count_user_request_total']['Павлов Андрей'], 1)   # IFE не считается

    def test_second_line_count_saved(self):
        self.assertEqual(self.stat.redirect_second_line_count, 1)
        self.assertEqual(self.stat.redirect_second_line, 33.3)

    def test_all_tickets_saved(self):
        self.assertEqual(self.stat.tickets.count(), 5)
        t = self.stat.tickets.get(ticket_id=2)
        self.assertEqual(t.category, 'Категория не выбрана')
        self.assertIsNone(t.first_answer_minutes)
        self.assertEqual(self.stat.tickets.get(ticket_id=1).first_answer_minutes, 30)

    def test_first_answer_rule(self):
        self.assertEqual(first_answer_minutes({'Date': '01.09.2026 10:00', 'fStartTime': '01.09.2026 10:05'}), 5)
        self.assertIsNone(first_answer_minutes({'Date': '01.09.2026 10:00', 'fStartTime': ''}))
        self.assertIsNone(first_answer_minutes({'Date': 'мусор', 'fStartTime': '01.09.2026 10:05'}))

    def test_every_report_downloads(self):
        for slug in REPORTS:
            self._download(slug)

    def test_crew_tablet_report(self):
        wb = self._download('crew_tablet')
        self.assertEqual(wb.sheetnames, ['ПАО "Аэрофлот"', 'АК "Россия"'])
        self.assertEqual(self._ids(wb['ПАО "Аэрофлот"']), [1, 2])
        self.assertEqual(self._ids(wb['АК "Россия"']), [10])
        self.assertEqual(wb['ПАО "Аэрофлот"']['D5'].hyperlink.target,
                         'https://helpdesk-su.ag-ife.com/request/1')

    def test_services_report(self):
        wb = self._download('services')
        self.assertEqual(self._ids(wb['ПАО "Аэрофлот"']), [3])
        self.assertEqual(self._ids(wb['АК "Россия"']), [11])

    def test_second_line_report(self):
        wb = self._download('second_line')
        self.assertEqual(self._ids(wb['ПАО "Аэрофлот"']), [2])
        self.assertEqual(self._ids(wb['АК "Россия"']), [])      # rportal не считается

    def test_admins_report_uses_saved_admins(self):
        wb = self._download('admins')
        self.assertEqual(self._ids(wb['ПАО "Аэрофлот"']), [1, 2])

    def test_first_answer_report_has_average_formula(self):
        ws = self._download('first_answer')['ПАО "Аэрофлот"']
        formulas = [c.value for c in ws['G'] if isinstance(c.value, str) and c.value.startswith('=')]
        self.assertEqual(len(formulas), 1)
        self.assertIn('AVERAGE(G5:G6)', formulas[0])

    def test_categories_report_single_sheet(self):
        wb = self._download('categories_akr')
        self.assertEqual(wb.sheetnames, ['АК "Россия"'])
        self.assertEqual(self._ids(wb['АК "Россия"']), [10])

    def test_unknown_report_404(self):
        response = self.client.get(reverse('stats_hd_export') + '?period=2026-09&report=nope')
        self.assertEqual(response.status_code, 404)

    def test_bad_period_400(self):
        response = self.client.get(reverse('stats_hd_export') + '?period=2026-13')
        self.assertEqual(response.status_code, 400)

    def test_old_period_without_tickets_shows_note(self):
        self.stat.tickets.all().delete()
        ws = self._download('crew_tablet')['ПАО "Аэрофлот"']
        self.assertIn('Обновить статистику', ws['C5'].value)


    def test_dashboard_has_export_buttons_and_json(self):
        response = self.client.get(reverse('stats_hd') + '?period=2026-09')
        content = response.content.decode()
        for slug in REPORTS:
            if slug != 'summary':
                self.assertIn(f'report={slug}', content)
        self.assertIn('id="dashboard-data"', content)

    def test_dashboard_bad_period_falls_back(self):
        response = self.client.get(reverse('stats_hd') + '?period=abc')
        self.assertEqual(response.status_code, 200)

    def test_table_starts_at_c3_centered_without_subtitle(self):
        ws = self._download('admins')['ПАО "Аэрофлот"']
        self.assertIsNone(ws['A1'].value)
        self.assertEqual(ws['C3'].value, 'ПАО "Аэрофлот"')
        self.assertEqual(ws['C4'].value, '№')
        for c in ws[5]:
            if c.value is not None:
                self.assertEqual(c.alignment.horizontal, 'center', c.coordinate)
        self.assertTrue(ws.print_options.horizontalCentered)

    def test_no_freeze_panes_and_filters(self):
        for slug in REPORTS:
            for ws in self._download(slug):
                self.assertIsNone(ws.freeze_panes, f'{slug}/{ws.title}')
                self.assertIsNone(ws.auto_filter.ref, f'{slug}/{ws.title}')

    def test_summary_categories_one_row_per_ticket(self):
        wb = self._download('summary')
        ws = wb['Категории Аэрофлот']
        self.assertEqual([c.value for c in ws[4] if c.value],
                         ['№', 'Дата создания', 'Категория', 'Исполнитель', 'Ссылка на заявку'])
        links = [c.value for c in ws['G'] if isinstance(c.value, str) and '/request/' in c.value]
        # равные по количеству категории — по алфавиту: «Категория не выбрана» (#2) раньше «Отчёты» (#1)
        self.assertEqual(links, ['https://helpdesk-su.ag-ife.com/request/2',
                                 'https://helpdesk-su.ag-ife.com/request/1'])
        self.assertEqual(ws['E5'].value, 'Категория не выбрана')
        self.assertEqual(ws['F5'].value, 'Тараторенков Юрий')
        ws_akr = wb['Категории Россия']
        self.assertEqual(ws_akr['G5'].value, 'https://helpdesk-fv.ag-ife.com/request/10')

    def test_dashboard_admins_have_ticket_links(self):
        response = self.client.get(reverse('stats_hd') + '?period=2026-09')
        admins = {a['name']: a for a in response.context['dashboard_data']['stat']['admins']}
        self.assertEqual([t['id'] for t in admins['Павлов Андрей']['tickets']], [1])
        self.assertEqual(admins['Романов Владислав']['tickets'],
                         [{'id': 10, 'url': 'https://helpdesk-fv.ag-ife.com/request/10'}])